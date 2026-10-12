"""Build an editable local application package from its fixed evidence snapshot."""
from __future__ import annotations

import json
import os
import re
import tempfile
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt


if TYPE_CHECKING:
    from careeros.applications.service import Application
    from careeros.jobs.models import JobListing


class ApplicationPackageExporter:
    def _document(self, text: str, output: Path) -> None:
        document = Document()
        # The upstream blank template carries an unrelated thumbnail attachment.
        # Drop that relationship so the package contains no images at all.
        package = document.part.package
        for identifier, relationship in list(package.rels.items()):
            if relationship.reltype.endswith('/thumbnail'):
                del package.rels[identifier]
        document.styles['Normal'].font.name = 'Arial'
        document.styles['Normal'].font.size = Pt(11)
        # Text is inserted literally, never interpreted as XML/HTML or fields.
        for line in text.splitlines():
            heading = re.fullmatch(r'#{1,3} (.+)', line)
            if heading:
                document.add_heading(heading[1], level=min(3, len(line) - len(line.lstrip('#'))))
            else:
                document.add_paragraph(line)
        document.save(output)

    def export(self, application: Application, job: JobListing, output: Path) -> Path:
        output = Path(output)
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.application-', dir=output.parent) as temporary:
            folder = Path(temporary)
            self._document(application.draft['resume_text'], folder / 'resume.docx')
            self._document(application.draft['cover_letter'], folder / 'cover-letter.docx')
            answers = '\n\n'.join(f'## {question}\n\n{answer}' for question, answer in application.draft['answers'].items())
            (folder / 'application-answers.md').write_text(answers + '\n', encoding='utf-8')
            source = {'application_id': application.id, 'snapshot_hash': application.snapshot_hash,
                      'draft_version': application.draft_version, 'job': application.snapshot['job']}
            (folder / 'job-source.json').write_text(json.dumps(source, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            archive = folder / 'package.zip'
            with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as package:
                for name in ('resume.docx', 'cover-letter.docx', 'application-answers.md', 'job-source.json'):
                    package.write(folder / name, name)
            os.replace(archive, output)
        return output
