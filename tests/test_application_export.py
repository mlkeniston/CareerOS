from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import json
import zipfile

import pytest

from test_applications_api import application_inputs


@pytest.fixture
def application_with_clearance(client):
    inputs = application_inputs(client)
    started = client.post('/api/applications', json=inputs).json()
    client.app.state.worker.run_once()
    client.app.state.worker.run_once()
    service = client.app.state.application_service
    application = service.get(started['id'])
    from careeros.jobs.verification import VerificationResult
    saved = application.snapshot['job']['verification']
    verifier = SimpleNamespace(calls=[], result=VerificationResult(**{k: saved[k] for k in ('status', 'reason', 'final_url', 'checked_at', 'text', 'fingerprint')}))
    def verify(url, company, title):
        verifier.calls.append((url, company, title))
        return verifier.result
    verifier.verify = verify
    service.posting_verifier = verifier
    route = f"/api/applications/{application.id}"
    return SimpleNamespace(client=client, service=service, id=application.id, verifier=verifier, route=route,
                           approve=lambda: client.post(route + '/approve'), export=lambda: client.get(route + '/export'))


def test_export_contains_ats_resume_letter_answers_and_source(application_with_clearance):
    from docx import Document
    p = application_with_clearance
    assert p.approve().status_code == 200
    response = p.export()
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/zip'
    archive = zipfile.ZipFile(BytesIO(response.content))
    assert set(archive.namelist()) == {'resume.docx', 'cover-letter.docx', 'application-answers.md', 'job-source.json'}
    for name in ('resume.docx', 'cover-letter.docx'):
        document = Document(BytesIO(archive.read(name)))
        assert 'Built a Python API' in '\n'.join(p.text for p in document.paragraphs)
        assert document.tables == []
        assert len(document.inline_shapes) == 0
        xml = document.element.xml
        assert '<w:vanish' not in xml
        assert 'w:num="2"' not in xml
        assert '<w:drawing' not in xml
        internal = zipfile.ZipFile(BytesIO(archive.read(name)))
        assert not any(part.endswith(('.jpeg', '.png', '.jpg')) for part in internal.namelist())
    source = json.loads(archive.read('job-source.json'))
    assert source['snapshot_hash'] == p.service.get(p.id).snapshot_hash
    assert source['job']['company'] == 'Acme'
    assert 'resume' not in source and 'profile' not in source
    with p.service.db.connect() as conn:
        record = conn.execute('SELECT * FROM application_exports WHERE application_id=?', (p.id,)).fetchone()
    assert Path(record['path']).read_bytes() == response.content
    assert record['exported_at']
    assert Path(record['path']).name == f'application-{p.id}.zip'
    assert list(Path(record['path']).parent.iterdir()) == [Path(record['path'])]
    assert len(p.verifier.calls) == 2


@pytest.mark.parametrize('state', ['needs_review', 'manager_review', 'quality_hold'])
def test_export_requires_explicit_approval(application_with_clearance, state):
    p = application_with_clearance
    with p.service.db.connect() as conn:
        conn.execute('UPDATE applications SET status=? WHERE id=?', (state, p.id))
    assert p.export().status_code == 409
    assert not (p.service.db.path.parent / 'exports' / f'application-{p.id}.zip').exists()


@pytest.mark.parametrize('action', ['approve', 'export'])
@pytest.mark.parametrize('change', ['text', 'fingerprint', 'final_url', 'held', 'exception'])
def test_fresh_fetch_failure_holds_without_refreshing_snapshot(application_with_clearance, action, change):
    p = application_with_clearance
    if action == 'export':
        assert p.approve().status_code == 200
    before = p.service.get(p.id)
    if change == 'exception':
        def unavailable(*args):
            raise OSError('offline')
        p.verifier.verify = unavailable
    else:
        changes = {'text': {'text': p.verifier.result.text + ' Changed'}, 'fingerprint': {'fingerprint': 'new'},
                   'final_url': {'final_url': 'https://jobs.example/redirect'}, 'held': {'status': 'held', 'reason': 'Page closed'}}
        p.verifier.result = replace(p.verifier.result, **changes[change])
    response = getattr(p, action)()
    assert response.status_code == 409
    after = p.service.get(p.id)
    assert after.status == 'quality_hold'
    assert after.hold_reason
    assert after.snapshot == before.snapshot
    assert after.snapshot_hash == before.snapshot_hash
    with p.service.db.connect() as conn:
        assert conn.execute('SELECT COUNT(*) FROM application_exports').fetchone()[0] == 0
    assert not (p.service.db.path.parent / 'exports' / f'application-{p.id}.zip').exists()


def test_stale_clearance_and_invalid_claims_block_package(application_with_clearance):
    p = application_with_clearance
    assert p.approve().status_code == 200
    with p.service.db.connect() as conn:
        conn.execute('UPDATE applications SET clearance_stale=1 WHERE id=?', (p.id,))
    assert p.export().status_code == 409
    with p.service.db.connect() as conn:
        conn.execute('UPDATE applications SET clearance_stale=0 WHERE id=?', (p.id,))
        draft = p.service.get(p.id).draft
        draft['resume_text'] = 'Built an AWS cluster'
        conn.execute('UPDATE applications SET draft_json=? WHERE id=?', (json.dumps(draft), p.id))
    assert p.export().status_code == 409


def test_export_failure_cleans_temporary_files_and_does_not_record_path(application_with_clearance, monkeypatch):
    from careeros.applications.export import ApplicationPackageExporter
    p = application_with_clearance
    assert p.approve().status_code == 200
    def fail(*args):
        raise OSError('disk full')
    monkeypatch.setattr(ApplicationPackageExporter, '_document', fail)
    with pytest.raises(OSError, match='disk full'):
        p.service.export_package(p.id)
    folder = p.service.db.path.parent / 'exports'
    assert not folder.exists() or not list(folder.iterdir())
    with p.service.db.connect() as conn:
        assert conn.execute('SELECT COUNT(*) FROM application_exports').fetchone()[0] == 0


def test_export_rechecks_concurrent_edit_before_writing_final_file(application_with_clearance, monkeypatch):
    from careeros.applications.export import ApplicationPackageExporter
    p = application_with_clearance
    assert p.approve().status_code == 200
    original = ApplicationPackageExporter.export
    def concurrent(exporter, application, job, output):
        path = original(exporter, application, job, output)
        with p.service.db.connect() as conn:
            conn.execute("UPDATE applications SET status='manager_review',clearance_stale=1,draft_version=draft_version+1 WHERE id=?", (p.id,))
        return path
    monkeypatch.setattr(ApplicationPackageExporter, 'export', concurrent)
    assert p.export().status_code == 409
    assert not (p.service.db.path.parent / 'exports' / f'application-{p.id}.zip').exists()
    with p.service.db.connect() as conn:
        assert conn.execute('SELECT COUNT(*) FROM application_exports').fetchone()[0] == 0


def test_fresh_manager_fetch_failure_is_a_persisted_blocking_review(application_with_clearance):
    p = application_with_clearance
    from careeros.applications.models import ApplicationDraft
    application = p.service.get(p.id)
    p.service.edit(p.id, ApplicationDraft.parse(application.draft))
    p.verifier.result = replace(p.verifier.result, status='held', reason='Employer page redirects')
    calls = len(p.client.app.state.fake_planner.calls)
    p.client.app.state.worker.run_once()
    assert p.service.get(p.id).status == 'quality_hold'
    assert 'redirects' in p.service.get(p.id).hold_reason
    assert len(p.client.app.state.fake_planner.calls) == calls
    review = p.client.app.state.hiring_manager_service.history(p.id)[-1]
    assert review['verdict'] == 'revise'
    assert review['findings'][-1]['severity'] == 'blocking'


def test_export_during_posting_change_persists_hold_and_cleans_build(application_with_clearance, monkeypatch):
    from careeros.applications.export import ApplicationPackageExporter
    p = application_with_clearance
    assert p.approve().status_code == 200
    original = ApplicationPackageExporter.export
    def changing(exporter, application, job, output):
        path = original(exporter, application, job, output)
        with p.service.db.connect() as conn:
            conn.execute("UPDATE job_verifications SET fingerprint='changed' WHERE job_id=?", (application.job_id,))
        return path
    monkeypatch.setattr(ApplicationPackageExporter, 'export', changing)
    assert p.export().status_code == 409
    assert p.service.get(p.id).status == 'quality_hold'
    assert p.service.get(p.id).hold_reason
    assert not list((p.service.db.path.parent / 'exports').iterdir())


def test_three_unresolved_reviews_have_a_persisted_hold_reason(application_with_clearance):
    from careeros.planner.protocols import PlannerResult
    from test_hiring_manager import BLOCK
    p = application_with_clearance
    p.service.reject(p.id, 'Recheck emphasis')
    p.client.app.state.worker.run_once()
    planner = p.client.app.state.fake_planner
    planner.result = PlannerResult(json.dumps(BLOCK))
    p.client.app.state.worker.run_once()
    planner.result = None
    p.client.app.state.worker.run_once()
    planner.result = PlannerResult(json.dumps(BLOCK))
    p.client.app.state.worker.run_once()
    application = p.service.get(p.id)
    assert application.status == 'quality_hold'
    assert 'three reviews' in application.hold_reason
    assert 'Weak emphasis' in application.hold_reason


def test_failed_export_record_restores_a_previous_complete_package(application_with_clearance):
    import sqlite3
    p = application_with_clearance
    assert p.approve().status_code == 200
    assert p.export().status_code == 200
    final = p.service.db.path.parent / 'exports' / f'application-{p.id}.zip'
    original = final.read_bytes()
    with p.service.db.connect() as conn:
        conn.execute("CREATE TRIGGER fail_export_record BEFORE INSERT ON application_exports BEGIN SELECT RAISE(ABORT,'failed export record'); END")
    with pytest.raises(sqlite3.IntegrityError, match='failed export record'):
        p.service.export_package(p.id)
    assert final.read_bytes() == original
    assert list(final.parent.iterdir()) == [final]


def test_approval_revalidates_claims_even_for_already_approved_application(application_with_clearance):
    p = application_with_clearance
    assert p.approve().status_code == 200
    draft = p.service.get(p.id).draft
    draft['cover_letter'] = 'Built an AWS cluster'
    with p.service.db.connect() as conn:
        conn.execute('UPDATE applications SET draft_json=? WHERE id=?', (json.dumps(draft), p.id))
    assert p.approve().status_code == 409
